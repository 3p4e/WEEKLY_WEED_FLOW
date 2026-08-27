# Offline proof of the formatting core: the canonical merged engine builds
# real bilingual documents and the PASS gate actually gates.
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import builder  # noqa: E402

SOP_MD = """<!--HEADERDATA
mk_title: Тест СОП за примерок
en_title: Test Sampling SOP
code: QCSOP-999
version: 1.0
doctype: SOP
orient: portrait
-->
# 1.0 ЦЕЛ|PURPOSE
Оваа процедура ја опишува постапката за земање мостри.|This procedure describes sampling.

# 3.0 ОДГОВОРНОСТИ|RESPONSIBILITIES
[[TABLE]]
Улога~~Role|||Одговорност~~Responsibility
Аналитичар~~Analyst|||Земање мостра~~Sampling
[[/TABLE]]

# 6.0 ПОСТАПКА|PROCEDURE
Чекор еден: подготовка.|Step one: preparation.
"""

FORM_MD = """<!--HEADERDATA
mk_title: Образец за земање мостра
en_title: Sampling Form
code: FM-QC-999
version: 1.0
doctype: FORM
orient: portrait
-->
# 1.0 ИНФОРМАЦИИ|INFORMATION
[[FORM:grid]]
Датум~~Date|||
Серија~~Batch|||
Материјал~~Material|||
Аналитичар~~Analyst|||
[[/FORM]]
"""


def test_sop_builds_and_passes(tmp_path):
    r = builder.build(SOP_MD, tmp_path, "sop_test")
    assert r.path.exists() and r.bytes > 10_000
    assert "RESULT: PASS" in r.verify_report
    assert r.doctype == "SOP"


def test_form_grid_builds_and_passes(tmp_path):
    r = builder.build(FORM_MD, tmp_path, "form_test")
    assert r.path.exists() and r.bytes > 10_000
    assert "RESULT: PASS" in r.verify_report
    assert r.doctype == "FORM"


def test_verify_gate_blocks_and_deletes(tmp_path, monkeypatch):
    # Force the verifier to fail: the artifact must NOT survive on disk.
    monkeypatch.setattr(builder, "run_verify",
                         lambda p, min_pt=6.0, require_bilingual=True: (False, "RESULT: FAIL"))
    with pytest.raises(builder.VerifyFailed) as e:
        builder.build(SOP_MD, tmp_path, "gated")
    assert "FAIL" in e.value.report
    assert not (tmp_path / "gated.docx").exists()


def _fake_engine_build(text, min_pt=10):
    """Bypass the real house-style engine and write a minimal, controllable
    .docx directly — isolates exactly one verify dimension per test instead
    of depending on the real template's own boilerplate content."""
    import docx as docxlib

    def _write(src, out):
        d = docxlib.Document()
        p = d.add_paragraph(text)
        p.runs[0].font.size = docxlib.shared.Pt(min_pt)
        d.save(out)
    return _write


def test_bilingual_missing_now_fails_the_gate(tmp_path, monkeypatch):
    # B1: the bilingual check used to only print a WARN and never gate the
    # build. An English-only document must now be refused.
    monkeypatch.setattr(builder.build_from_md, "main",
                         _fake_engine_build("Only English content appears anywhere in this "
                                            "document body, repeated for length. " * 20))
    with pytest.raises(builder.VerifyFailed) as e:
        builder.build(SOP_MD, tmp_path, "bilingual_fail")
    assert "FAIL (missing a language)" in e.value.report


def test_bilingual_optout_allows_monolingual_document(tmp_path, monkeypatch):
    # The forward-compatible `bilingual: no` HEADERDATA escape hatch.
    monkeypatch.setattr(builder.build_from_md, "main",
                         _fake_engine_build("Only English content appears anywhere in this "
                                            "document body, repeated for length. " * 20))
    md = SOP_MD.replace("orient: portrait\n-->", "orient: portrait\nbilingual: no\n-->")
    r = builder.build(md, tmp_path, "bilingual_optout")
    assert "RESULT: PASS" in r.verify_report


def test_fidelity_shortfall_fails_the_gate(tmp_path, monkeypatch):
    # B1: a produced document that is a content SHORTFALL vs. its own source
    # markdown must be refused — the impoverishment/fabrication-by-omission
    # class of failure.
    monkeypatch.setattr(builder.build_from_md, "main", _fake_engine_build("Кратко. Short."))
    with pytest.raises(builder.VerifyFailed) as e:
        builder.build(SOP_MD, tmp_path, "fidelity_fail")
    assert "FIDELITY" in e.value.report and "FAIL" in e.value.report


def test_parse_survives_embedded_comment_terminator_in_value():
    # B2: a HEADERDATA field value containing a literal '-->' must not be
    # mistaken for the block's own terminator.
    import build_from_md
    md = (
        "<!--HEADERDATA\n"
        "mk_title: Опис --> на нешто\n"
        "en_title: Description\n"
        "code: C-1\n"
        "version: 1.0\n"
        "doctype: SOP\n"
        "orient: portrait\n"
        "-->\n"
        "# 1.0 ЦЕЛ|PURPOSE\n"
        "Текст.|Text.\n"
    )
    hd, blocks = build_from_md.parse(md)
    assert hd["mk_title"] == "Опис --> на нешто"
    assert hd["en_title"] == "Description"
    assert hd["code"] == "C-1"
    assert hd["doctype"] == "SOP"
    body_text = " ".join(str(b) for b in blocks)
    assert "en_title" not in body_text and "doctype" not in body_text and "version" not in body_text


def test_parse_rejects_duplicate_headerdata_key():
    # HIGH: last-line-wins on a repeated key is exactly what would let an
    # injected/forged `code:` line (smuggled in via an embedded newline in
    # some upstream field, e.g. meta.orient) silently override the real
    # code — diverging the built .docx's printed identity from the
    # audit-trail registry row. Reject at parse time too, independent of
    # whether the caller that assembled the Markdown sanitized its inputs.
    import build_from_md
    md = (
        "<!--HEADERDATA\n"
        "mk_title: МК\n"
        "en_title: EN\n"
        "code: SOP-REAL\n"
        "code: SOP-FORGED\n"
        "version: 1.0\n"
        "doctype: SOP\n"
        "orient: portrait\n"
        "-->\n"
        "# 1.0 ЦЕЛ|PURPOSE\n"
        "Текст.|Text.\n"
    )
    with pytest.raises(ValueError):
        build_from_md.parse(md)


def test_house_style_in_output(tmp_path):
    # The produced docx carries the house navy #2B547E and Calibri.
    import zipfile
    r = builder.build(FORM_MD, tmp_path, "style_test")
    xml = zipfile.ZipFile(r.path).read("word/document.xml").decode("utf-8")
    assert "2B547E" in xml
    assert "Calibri" in xml


def test_safe_name():
    assert builder.safe_name("../../etc/passwd") == ".._.._etc_passwd"
    assert builder.safe_name("") == "document"
    assert "/" not in builder.safe_name("a/b\\c d")


def test_concurrent_builds_of_same_code_produce_two_intact_documents(tmp_path):
    """H13 — the output path used to be derived from the document CODE alone,
    so two concurrent builds of the same code wrote the same file: they
    interleaved, and the FAIL cleanup could delete the OTHER build's passing
    document. The in-process lock cannot help, because the service runs
    multiple uvicorn worker PROCESSES.

    Each build now owns a unique path, so both artifacts survive intact."""
    import concurrent.futures

    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as ex:
        results = [f.result() for f in
                   [ex.submit(builder.build, SOP_MD, tmp_path, "same_code") for _ in range(4)]]

    paths = [r.path for r in results]
    assert len({str(p) for p in paths}) == 4, "builds collided on one path"
    for r in results:
        assert r.path.exists() and r.path.stat().st_size == r.bytes > 0
        assert r.path.name.startswith("same_code-") and r.path.suffix == ".docx"
    # nothing half-written is left behind under the dot-prefixed staging name
    assert not list(tmp_path.glob(".*partial*"))


def test_new_sop_raises_loudly_when_template_missing(tmp_path):
    """MEDIUM: new_sop(from_template=True) used to silently fall back to a
    bare Document() with NO mandatory header/footer/logo when PP_TEMPLATE
    didn't exist -- no log, no exception -- and that bare document could still
    PASS pp_verify's gate (it never checks the header/footer table exists).
    A missing template file (bad container image, bad relative path) must
    fail loudly instead."""
    import pp_format
    missing = tmp_path / "does_not_exist.docx"
    with pytest.raises(FileNotFoundError, match=str(missing)):
        pp_format.new_sop(code="C-1", template=str(missing))


def test_new_annex_raises_loudly_when_template_missing(tmp_path):
    import pp_format
    missing = tmp_path / "does_not_exist.docx"
    with pytest.raises(FileNotFoundError, match=str(missing)):
        pp_format.new_annex(code="C-1", template=str(missing))


def test_new_sop_from_template_false_is_unaffected():
    """The other branch must keep working: a caller that deliberately asks
    for NO template (from_template=False) still gets a plain Document(), not
    an exception -- the guard is scoped to the "wanted a template but it's
    missing" case only."""
    import pp_format
    d = pp_format.new_sop(from_template=False)
    assert d is not None


def test_new_annex_from_template_false_is_unaffected():
    import pp_format
    d = pp_format.new_annex(from_template=False)
    assert d is not None


def test_new_sop_succeeds_with_the_real_template():
    """The happy path is untouched by the guard: the real PP_TEMPLATE exists
    in this checkout, so from_template=True (the default) must still build."""
    import pp_format
    d = pp_format.new_sop(code="C-1", mk_title="Наслов", en_title="Title")
    assert d is not None


def test_apply_pp_header_raises_when_header_table_missing():
    """BUG 6 — apply_pp_header used to `return False` on a missing header
    table, and both new_sop/new_annex ignore that return value, so a base
    template whose header structure had drifted would silently ship a
    document with a blank Document-name/Code/Version field. Match the same
    "loud failure over silent fallback" precedent already applied to the
    missing-template-file case (see test_new_sop_raises_loudly_when_template_missing)."""
    import pp_format
    from docx import Document
    d = Document()  # no header table at all
    with pytest.raises(RuntimeError, match="header table not found"):
        pp_format.apply_pp_header(d, "MK", "CODE-1", "EN")


def test_apply_pp_header_raises_on_run_index_mismatch():
    """BUG 6 — the setrun() helper used to silently no-op whenever the
    expected run index wasn't present (structure drifted), leaving that one
    header field blank with no error anywhere. It must now raise instead."""
    import pp_format
    from docx import Document
    from docx.shared import Cm
    d = Document()
    h = d.sections[0].header
    h.is_linked_to_previous = False
    t = h.add_table(rows=1, cols=3, width=Cm(18))
    # cell(0,1).paragraphs[1] is expected to carry >= 6 runs; give it 1.
    p1 = t.cell(0, 1).add_paragraph()
    p1.add_run("only one run")
    with pytest.raises(RuntimeError, match="document-name cell"):
        pp_format.apply_pp_header(d, "MK title", "CODE-1", "EN title")


def test_new_sop_succeeds_with_the_real_template_header_intact():
    """The real PP_TEMPLATE's header must still satisfy every index BUG 6's
    guard now enforces -- the happy path is unaffected by turning the
    silent no-op into a raise."""
    import pp_format
    d = pp_format.new_sop(code="C-1", mk_title="Наслов", en_title="Title", version="2.0")
    h = d.sections[0].header.tables[0]
    nm = h.cell(0, 1).paragraphs[1].runs
    assert nm[0].text == "Наслов "
    assert nm[5].text == "Title"
    cd = h.cell(0, 2).paragraphs[2].runs
    assert cd[0].text == "C-1"
    vr = h.cell(1, 2).paragraphs[0].runs
    assert vr[-1].text == "2.0"


def test_omath_count_does_not_double_count_oMathPara_wrapper(tmp_path):
    """BUG 8 — a bare doc.count("<m:oMath") also matched every <m:oMathPara>
    wrapper (a display equation is <m:oMathPara><m:oMath>...</m:oMath></m:oMathPara>),
    over-counting the equation total surfaced to API callers as verification
    evidence. One real display equation must report omath=1, not 2."""
    import zipfile
    import pp_verify

    src = tmp_path / "src.docx"
    from docx import Document
    Document().save(src)

    with zipfile.ZipFile(src) as zin:
        names = zin.namelist()
        doc_xml = zin.read("word/document.xml").decode("utf-8")
        others = {n: zin.read(n) for n in names if n != "word/document.xml"}

    # One display equation: <m:oMathPara> wrapping one <m:oMath>. A bare
    # substring count of "<m:oMath" would see 2 hits here; there is 1 equation.
    fragment = "<m:oMathPara><m:oMath><m:r><m:t>x</m:t></m:r></m:oMath></m:oMathPara>"
    doc_xml = doc_xml.replace("</w:body>", fragment + "</w:body>")

    out = tmp_path / "with_eqn.docx"
    with zipfile.ZipFile(out, "w") as zout:
        zout.writestr("word/document.xml", doc_xml)
        for n, b in others.items():
            zout.writestr(n, b)

    r = pp_verify.analyse(str(out))
    assert r["omath"] == 1, f"expected exactly 1 equation, got {r['omath']}"


def test_min_font_ignores_unanchored_text_that_resembles_a_sz_element():
    """ITEM 9 (investigated) — the font-floor regex used to be a bare
    substring match with no leading '<' anchor, so literal text containing
    "w:sz w:val=\"N\"" that is NOT a real <w:sz> element (e.g. embedded in a
    run's own text content) would be misread as a font-size declaration.
    Real generated output never produces this (OOXML always orders a
    border's w:val before its w:sz, and this codebase's own border helpers
    do too — see the code comment), but the regex itself must not depend on
    that alone."""
    import pp_verify
    xml = ('<w:t>note: w:sz w:val="2" is out of range</w:t>'
           '<w:rPr><w:sz w:val="16"/></w:rPr>')
    assert pp_verify._min_font({"word/document.xml": xml}) == 16


def test_min_font_still_finds_the_real_minimum_across_parts():
    import pp_verify
    parts = {
        "word/document.xml": '<w:rPr><w:sz w:val="22"/></w:rPr>',
        "word/header1.xml": '<w:rPr><w:sz w:val="14"/></w:rPr>',
        "word/footer1.xml": '<w:rPr><w:sz w:val="12"/></w:rPr>',
        "word/styles.xml": '<w:rPr><w:sz w:val="1"/></w:rPr>',  # not doc/header/footer -> ignored
    }
    assert pp_verify._min_font(parts) == 12


def test_failed_build_cannot_delete_another_builds_document(tmp_path, monkeypatch):
    """H13 — the FAIL path unlinks its own staging file, which is uniquely
    named, so it can no longer take out a sibling build's passing artifact."""
    good = builder.build(SOP_MD, tmp_path, "shared_code")
    assert good.path.exists()
    monkeypatch.setattr(builder, "run_verify", lambda *a, **k: (False, "RESULT: FAIL"))
    with pytest.raises(builder.VerifyFailed):
        builder.build(SOP_MD, tmp_path, "shared_code")     # same code, must FAIL
    assert good.path.exists(), "a failed build deleted a different build's document"
    assert good.path.stat().st_size == good.bytes
