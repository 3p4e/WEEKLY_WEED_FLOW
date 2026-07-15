#!/usr/bin/env python3
"""
Excel Integration Tool for Cannabis EU GMP QMS Creator
Converts facility_data.yaml ↔ Excel for easy editing by non-technical users
"""

import sys
from pathlib import Path
from typing import Dict, List, Any
import yaml
import pandas as pd
from datetime import datetime
from rich.console import Console
from rich.table import Table

console = Console()


class ExcelIntegration:
    """Convert between YAML configuration and Excel spreadsheet."""

    def __init__(self, config_path: Path, excel_path: Path):
        """Initialize Excel integration."""
        self.config_path = config_path
        self.excel_path = excel_path
        self.data: Dict[str, Any] = {}

    def yaml_to_excel(self) -> bool:
        """Convert facility_data.yaml to Excel workbook."""
        console.print("\n[bold cyan]Converting YAML → Excel[/bold cyan]")
        console.print("=" * 80 + "\n")

        try:
            # Load YAML data
            with open(self.config_path, 'r', encoding='utf-8') as f:
                self.data = yaml.safe_load(f)

            console.print("[cyan]Creating Excel workbook with sheets...[/cyan]")

            # Create Excel writer
            with pd.ExcelWriter(self.excel_path, engine='openpyxl') as writer:

                # Sheet 1: Company Information
                df_company = self._create_company_sheet()
                df_company.to_excel(writer, sheet_name='Company', index=False)
                console.print("  ✓ Company sheet created")

                # Sheet 2: Regulatory
                df_regulatory = self._create_regulatory_sheet()
                df_regulatory.to_excel(writer, sheet_name='Regulatory', index=False)
                console.print("  ✓ Regulatory sheet created")

                # Sheet 3: Personnel
                df_personnel = self._create_personnel_sheet()
                df_personnel.to_excel(writer, sheet_name='Personnel', index=False)
                console.print("  ✓ Personnel sheet created")

                # Sheet 4: Equipment
                df_equipment = self._create_equipment_sheet()
                df_equipment.to_excel(writer, sheet_name='Equipment', index=False)
                console.print("  ✓ Equipment sheet created")

                # Sheet 5: Facilities
                df_facilities = self._create_facilities_sheet()
                df_facilities.to_excel(writer, sheet_name='Facilities', index=False)
                console.print("  ✓ Facilities sheet created")

                # Sheet 6: Operations
                df_operations = self._create_operations_sheet()
                df_operations.to_excel(writer, sheet_name='Operations', index=False)
                console.print("  ✓ Operations sheet created")

                # Sheet 7: Quality Control
                df_qc = self._create_qc_sheet()
                df_qc.to_excel(writer, sheet_name='Quality Control', index=False)
                console.print("  ✓ Quality Control sheet created")

                # Sheet 8: Document Control
                df_doc = self._create_doc_control_sheet()
                df_doc.to_excel(writer, sheet_name='Document Control', index=False)
                console.print("  ✓ Document Control sheet created")

                # Sheet 9: Instructions
                df_instructions = self._create_instructions_sheet()
                df_instructions.to_excel(writer, sheet_name='INSTRUCTIONS', index=False)
                console.print("  ✓ Instructions sheet created")

            # Apply formatting
            self._format_excel_workbook()

            console.print(f"\n[bold green]✓ Excel file created successfully![/bold green]")
            console.print(f"[cyan]Location: {self.excel_path}[/cyan]\n")
            return True

        except Exception as e:
            console.print(f"[bold red]✗ Error: {e}[/bold red]")
            return False

    def excel_to_yaml(self) -> bool:
        """Convert Excel workbook back to facility_data.yaml."""
        console.print("\n[bold cyan]Converting Excel → YAML[/bold cyan]")
        console.print("=" * 80 + "\n")

        try:
            # Read all sheets
            excel_data = pd.read_excel(self.excel_path, sheet_name=None)

            # Build YAML structure
            yaml_data = {}

            # Company
            yaml_data['company'] = self._parse_company_sheet(excel_data['Company'])
            console.print("  ✓ Company data parsed")

            # Regulatory
            yaml_data['regulatory'] = self._parse_regulatory_sheet(excel_data['Regulatory'])
            console.print("  ✓ Regulatory data parsed")

            # Personnel
            yaml_data['personnel'] = self._parse_personnel_sheet(excel_data['Personnel'])
            console.print("  ✓ Personnel data parsed")

            # Equipment
            yaml_data['equipment'] = self._parse_equipment_sheet(excel_data['Equipment'])
            console.print("  ✓ Equipment data parsed")

            # Facilities
            yaml_data['facilities'] = self._parse_facilities_sheet(excel_data['Facilities'])
            console.print("  ✓ Facilities data parsed")

            # Operations
            yaml_data['operations'] = self._parse_operations_sheet(excel_data['Operations'])
            console.print("  ✓ Operations data parsed")

            # Quality Control
            yaml_data['quality_control'] = self._parse_qc_sheet(excel_data['Quality Control'])
            console.print("  ✓ Quality Control data parsed")

            # Document Control
            yaml_data['document_control'] = self._parse_doc_control_sheet(excel_data['Document Control'])
            console.print("  ✓ Document Control data parsed")

            # Add metadata
            yaml_data['metadata'] = {
                'qms_version': '1.0',
                'last_updated': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                'effective_date': '2026-02-01',
                'next_review_date': '2027-02-01',
                'generated_by': 'Cannabis EU GMP QMS Creator - Automation System',
                'generation_date': ''
            }

            # Write YAML file
            with open(self.config_path, 'w', encoding='utf-8') as f:
                yaml.dump(yaml_data, f, default_flow_style=False, allow_unicode=True, sort_keys=False)

            console.print(f"\n[bold green]✓ YAML file updated successfully![/bold green]")
            console.print(f"[cyan]Location: {self.config_path}[/cyan]\n")
            return True

        except Exception as e:
            console.print(f"[bold red]✗ Error: {e}[/bold red]")
            import traceback
            traceback.print_exc()
            return False

    # === Sheet Creation Methods ===

    def _create_company_sheet(self) -> pd.DataFrame:
        """Create Company sheet."""
        company = self.data.get('company', {})
        address = company.get('address', {})

        data = {
            'Field': [
                'Legal Entity Name',
                'Registration Number',
                'Facility Name',
                'Street Address',
                'City',
                'Region',
                'Postal Code',
                'Country',
                'Latitude',
                'Longitude'
            ],
            'Value': [
                company.get('legal_entity_name', ''),
                company.get('registration_number', ''),
                company.get('facility_name', ''),
                address.get('street', ''),
                address.get('city', ''),
                address.get('region', ''),
                address.get('postal_code', ''),
                address.get('country', ''),
                address.get('latitude', ''),
                address.get('longitude', '')
            ]
        }
        return pd.DataFrame(data)

    def _create_regulatory_sheet(self) -> pd.DataFrame:
        """Create Regulatory sheet."""
        regulatory = self.data.get('regulatory', {})
        license_info = regulatory.get('medical_cannabis_license', {})
        authority = regulatory.get('regulatory_authority', {})
        secondary = regulatory.get('secondary_authorities', {})
        moi = secondary.get('ministry_of_interior', {})
        lha = secondary.get('local_health_authority', {})

        data = {
            'Field': [
                '=== MEDICAL CANNABIS LICENSE ===',
                'License Number',
                'Issue Date',
                'Expiration Date',
                '',
                '=== PRIMARY REGULATORY AUTHORITY ===',
                'Authority Name',
                'Contact Name',
                'Phone',
                'Email',
                '',
                '=== MINISTRY OF INTERIOR ===',
                'Contact Name',
                'Phone',
                'Email',
                '',
                '=== LOCAL HEALTH AUTHORITY ===',
                'Contact Name',
                'Phone',
                'Email'
            ],
            'Value': [
                '',
                license_info.get('license_number', ''),
                license_info.get('issue_date', ''),
                license_info.get('expiration_date', ''),
                '',
                '',
                authority.get('name', ''),
                authority.get('contact_name', ''),
                authority.get('phone', ''),
                authority.get('email', ''),
                '',
                '',
                moi.get('contact_name', ''),
                moi.get('phone', ''),
                moi.get('email', ''),
                '',
                '',
                lha.get('contact_name', ''),
                lha.get('phone', ''),
                lha.get('email', '')
            ]
        }
        return pd.DataFrame(data)

    def _create_personnel_sheet(self) -> pd.DataFrame:
        """Create Personnel sheet."""
        personnel = self.data.get('personnel', {})

        roles = [
            ('qualified_person', 'Qualified Person (QP)'),
            ('facility_manager', 'Facility Manager'),
            ('qa_manager', 'QA Manager'),
            ('production_manager', 'Production Manager'),
            ('sanitation_manager', 'Sanitation Manager'),
            ('security_manager', 'Security Manager'),
            ('hr_manager', 'HR Manager'),
            ('records_manager', 'Records Manager')
        ]

        rows = []
        for role_key, role_name in roles:
            person = personnel.get(role_key, {})
            rows.append({
                'Role': role_name,
                'First Name': person.get('first_name', ''),
                'Last Name': person.get('last_name', ''),
                'Title': person.get('title', ''),
                'Qualifications': person.get('qualifications', ''),
                'Phone': person.get('phone', ''),
                'Email': person.get('email', '')
            })

        return pd.DataFrame(rows)

    def _create_equipment_sheet(self) -> pd.DataFrame:
        """Create Equipment sheet."""
        equipment = self.data.get('equipment', {})

        rows = []

        # Main equipment
        for eq_key in ['drying_machine', 'trimming_machine', 'packaging_equipment',
                       'hvac_system', 'irrigation_system', 'lighting_system',
                       'water_system', 'waste_system']:
            eq = equipment.get(eq_key, {})
            if eq:
                rows.append({
                    'Equipment ID': eq.get('equipment_id', ''),
                    'Name': eq.get('name', ''),
                    'Model': eq.get('model', eq.get('type', '')),
                    'Manufacturer': eq.get('manufacturer', ''),
                    'Serial Number': eq.get('serial_number', ''),
                    'Location': eq.get('location', ''),
                    'Calibration Due': eq.get('calibration_due_date', ''),
                    'Status': eq.get('qualification_status', '')
                })

        # Instruments
        instruments = equipment.get('instruments', {})
        for inst_key in ['instrument_1', 'instrument_2', 'instrument_3']:
            inst = instruments.get(inst_key, {})
            if inst:
                rows.append({
                    'Equipment ID': inst.get('equipment_id', ''),
                    'Name': inst.get('name', ''),
                    'Model': inst.get('model', ''),
                    'Manufacturer': '',
                    'Serial Number': inst.get('serial_number', ''),
                    'Location': '',
                    'Calibration Due': inst.get('calibration_due_date', ''),
                    'Status': ''
                })

        return pd.DataFrame(rows)

    def _create_facilities_sheet(self) -> pd.DataFrame:
        """Create Facilities sheet."""
        facilities = self.data.get('facilities', {})
        rooms = facilities.get('rooms', [])

        rows = []
        for room in rooms:
            rows.append({
                'Room ID': room.get('room_id', ''),
                'Room Name': room.get('name', ''),
                'Area (sqm)': room.get('area_sqm', ''),
                'Classification': room.get('classification', ''),
                'Purpose': room.get('purpose', '')
            })

        return pd.DataFrame(rows)

    def _create_operations_sheet(self) -> pd.DataFrame:
        """Create Operations sheet."""
        operations = self.data.get('operations', {})
        scope = operations.get('scope', {})
        objectives = operations.get('quality_objectives', {})
        metrics = operations.get('metrics', {})

        data = {
            'Field': [
                '=== SCOPE ===',
                'Cultivation',
                'Flowering',
                'Harvesting',
                'Drying & Curing',
                'Processing',
                'Packaging',
                'Quality Control',
                'Distribution',
                '',
                '=== ADDITIONAL INFO ===',
                'Additional Operations',
                'Excluded Operations',
                '',
                '=== QUALITY OBJECTIVES ===',
                'Objective 1',
                'Objective 2',
                'Objective 3',
                '',
                '=== METRICS ===',
                'Metric 1',
                'Metric 2',
                'Metric 3'
            ],
            'Value': [
                '',
                'Yes' if scope.get('cultivation') else 'No',
                'Yes' if scope.get('flowering') else 'No',
                'Yes' if scope.get('harvesting') else 'No',
                'Yes' if scope.get('drying_curing') else 'No',
                'Yes' if scope.get('processing') else 'No',
                'Yes' if scope.get('packaging') else 'No',
                'Yes' if scope.get('quality_control') else 'No',
                'Yes' if scope.get('distribution') else 'No',
                '',
                '',
                operations.get('additional_operations', ''),
                operations.get('excluded_operations', ''),
                '',
                '',
                objectives.get('objective_1', ''),
                objectives.get('objective_2', ''),
                objectives.get('objective_3', ''),
                '',
                '',
                metrics.get('target_metric_1', ''),
                metrics.get('target_metric_2', ''),
                metrics.get('target_metric_3', '')
            ]
        }
        return pd.DataFrame(data)

    def _create_qc_sheet(self) -> pd.DataFrame:
        """Create Quality Control sheet."""
        qc = self.data.get('quality_control', {})
        contract = qc.get('contract_lab_details', {})
        capabilities = qc.get('testing_capabilities', {})

        data = {
            'Field': [
                'In-House Lab',
                'Contract Lab',
                '',
                '=== CONTRACT LAB DETAILS ===',
                'Lab Name',
                'Address',
                'Contact Name',
                'Phone',
                'Email',
                'Accreditation',
                '',
                '=== TESTING CAPABILITIES ===',
                'Potency Testing',
                'Microbial Testing',
                'Pesticide Testing',
                'Heavy Metals Testing',
                'Moisture Content',
                'Visual Inspection'
            ],
            'Value': [
                'Yes' if qc.get('in_house_lab') else 'No',
                'Yes' if qc.get('contract_lab') else 'No',
                '',
                '',
                contract.get('name', ''),
                contract.get('address', ''),
                contract.get('contact_name', ''),
                contract.get('phone', ''),
                contract.get('email', ''),
                contract.get('accreditation', ''),
                '',
                '',
                capabilities.get('potency_testing', ''),
                capabilities.get('microbial_testing', ''),
                capabilities.get('pesticide_testing', ''),
                capabilities.get('heavy_metals_testing', ''),
                capabilities.get('moisture_content', ''),
                capabilities.get('visual_inspection', '')
            ]
        }
        return pd.DataFrame(data)

    def _create_doc_control_sheet(self) -> pd.DataFrame:
        """Create Document Control sheet."""
        doc = self.data.get('document_control', {})
        retention = doc.get('retention_periods', {})
        locations = doc.get('controlled_locations', [])
        access = doc.get('access_control', {})

        data = {
            'Field': [
                '=== RETENTION PERIODS ===',
                'Batch Records',
                'Training Records',
                'Equipment Qualification',
                'Deviation Reports',
                'Audit Reports',
                'QMS Documents',
                '',
                '=== CONTROLLED LOCATIONS ===',
                'Location 1',
                'Location 2',
                'Location 3',
                '',
                '=== ACCESS CONTROL ===',
                'Procedure',
                'Backup Frequency'
            ],
            'Value': [
                '',
                retention.get('batch_records', ''),
                retention.get('training_records', ''),
                retention.get('equipment_qualification', ''),
                retention.get('deviation_reports', ''),
                retention.get('audit_reports', ''),
                retention.get('qms_documents', ''),
                '',
                '',
                locations[0] if len(locations) > 0 else '',
                locations[1] if len(locations) > 1 else '',
                locations[2] if len(locations) > 2 else '',
                '',
                '',
                access.get('procedure', ''),
                access.get('backup_frequency', '')
            ]
        }
        return pd.DataFrame(data)

    def _create_instructions_sheet(self) -> pd.DataFrame:
        """Create instructions sheet."""
        data = {
            'INSTRUCTIONS FOR USING THIS EXCEL FILE': [
                '',
                '=== HOW TO USE ===',
                '1. Edit any values in the "Value" column across all sheets',
                '2. Do NOT modify the "Field" column or sheet names',
                '3. Do NOT delete or add rows (except in Personnel/Equipment if needed)',
                '4. Save the file when done editing',
                '5. Run: python scripts/excel_integration.py --import',
                '6. Your changes will be imported back to facility_data.yaml',
                '7. Run: python scripts/customize_documents.py --all',
                '8. All 100+ documents will be regenerated with your new data',
                '',
                '=== TIPS ===',
                '• Date format: YYYY-MM-DD (e.g., 2026-01-15)',
                '• Phone format: +389 70 123 456',
                '• Email format: name@domain.com',
                '• Leave cells blank if data is not available',
                '• Yes/No values for checkboxes',
                '',
                '=== SUPPORT ===',
                'For questions, refer to README_AUTOMATION.md',
                '',
                'Cannabis EU GMP QMS Creator - Automation System v1.0',
                'Purely Plant GmbH - Skopje, North Macedonia'
            ]
        }
        return pd.DataFrame(data)

    # === Sheet Parsing Methods ===

    def _parse_company_sheet(self, df: pd.DataFrame) -> Dict:
        """Parse Company sheet."""
        values = df.set_index('Field')['Value'].to_dict()
        return {
            'legal_entity_name': values.get('Legal Entity Name', ''),
            'registration_number': values.get('Registration Number', ''),
            'facility_name': values.get('Facility Name', ''),
            'address': {
                'street': values.get('Street Address', ''),
                'city': values.get('City', ''),
                'region': values.get('Region', ''),
                'postal_code': str(values.get('Postal Code', '')),
                'country': values.get('Country', ''),
                'latitude': str(values.get('Latitude', '')),
                'longitude': str(values.get('Longitude', ''))
            }
        }

    def _parse_regulatory_sheet(self, df: pd.DataFrame) -> Dict:
        """Parse Regulatory sheet."""
        values = df.set_index('Field')['Value'].to_dict()
        return {
            'medical_cannabis_license': {
                'license_number': values.get('License Number', ''),
                'issue_date': str(values.get('Issue Date', '')),
                'expiration_date': str(values.get('Expiration Date', ''))
            },
            'regulatory_authority': {
                'name': values.get('Authority Name', ''),
                'contact_name': values.get('Contact Name', ''),
                'phone': values.get('Phone', ''),
                'email': values.get('Email', '')
            },
            'secondary_authorities': {
                'ministry_of_interior': {
                    'contact_name': df[df['Field'] == 'Contact Name'].iloc[1]['Value'] if len(df[df['Field'] == 'Contact Name']) > 1 else '',
                    'phone': df[df['Field'] == 'Phone'].iloc[1]['Value'] if len(df[df['Field'] == 'Phone']) > 1 else '',
                    'email': df[df['Field'] == 'Email'].iloc[1]['Value'] if len(df[df['Field'] == 'Email']) > 1 else ''
                },
                'local_health_authority': {
                    'contact_name': df[df['Field'] == 'Contact Name'].iloc[2]['Value'] if len(df[df['Field'] == 'Contact Name']) > 2 else '',
                    'phone': df[df['Field'] == 'Phone'].iloc[2]['Value'] if len(df[df['Field'] == 'Phone']) > 2 else '',
                    'email': df[df['Field'] == 'Email'].iloc[2]['Value'] if len(df[df['Field'] == 'Email']) > 2 else ''
                }
            }
        }

    def _parse_personnel_sheet(self, df: pd.DataFrame) -> Dict:
        """Parse Personnel sheet."""
        personnel = {}

        role_mapping = {
            'Qualified Person (QP)': 'qualified_person',
            'Facility Manager': 'facility_manager',
            'QA Manager': 'qa_manager',
            'Production Manager': 'production_manager',
            'Sanitation Manager': 'sanitation_manager',
            'Security Manager': 'security_manager',
            'HR Manager': 'hr_manager',
            'Records Manager': 'records_manager'
        }

        for _, row in df.iterrows():
            role_display = row['Role']
            role_key = role_mapping.get(role_display)
            if role_key:
                personnel[role_key] = {
                    'first_name': row['First Name'],
                    'last_name': row['Last Name'],
                    'title': row['Title'],
                    'qualifications': row['Qualifications'],
                    'phone': row['Phone'],
                    'email': row['Email']
                }

        return personnel

    def _parse_equipment_sheet(self, df: pd.DataFrame) -> Dict:
        """Parse Equipment sheet."""
        equipment = {}
        instruments = {}

        for idx, row in df.iterrows():
            eq_id = row['Equipment ID']

            # Determine equipment key from ID
            if 'DRYING' in eq_id:
                key = 'drying_machine'
            elif 'TRIM' in eq_id:
                key = 'trimming_machine'
            elif 'PKG' in eq_id:
                key = 'packaging_equipment'
            elif 'HVAC' in eq_id:
                key = 'hvac_system'
            elif 'IRR' in eq_id:
                key = 'irrigation_system'
            elif 'LIGHT' in eq_id:
                key = 'lighting_system'
            elif 'WATER' in eq_id:
                key = 'water_system'
            elif 'WASTE' in eq_id:
                key = 'waste_system'
            elif 'QC' in eq_id:
                # Instrument
                inst_num = eq_id.split('-')[-1]
                key = f'instrument_{inst_num}'
                instruments[key] = {
                    'name': row['Name'],
                    'model': row['Model'],
                    'serial_number': row['Serial Number'],
                    'equipment_id': eq_id,
                    'calibration_due_date': str(row['Calibration Due'])
                }
                continue
            else:
                continue

            equipment[key] = {
                'name': row['Name'],
                'model': row['Model'],
                'manufacturer': row['Manufacturer'],
                'serial_number': row['Serial Number'],
                'equipment_id': eq_id,
                'location': row['Location'],
                'calibration_due_date': str(row['Calibration Due']),
                'qualification_status': row['Status']
            }

        if instruments:
            equipment['instruments'] = instruments

        return equipment

    def _parse_facilities_sheet(self, df: pd.DataFrame) -> Dict:
        """Parse Facilities sheet."""
        rooms = []

        for _, row in df.iterrows():
            rooms.append({
                'room_id': row['Room ID'],
                'name': row['Room Name'],
                'area_sqm': str(row['Area (sqm)']),
                'classification': row['Classification'],
                'purpose': row['Purpose']
            })

        return {
            'total_area_sqm': '2500',
            'rooms': rooms,
            'storage': {
                'primary_document_storage': 'Quality Assurance Office, Server Room',
                'backup_document_storage': 'Off-site Cloud Storage + Local Backup Server',
                'archive_storage': 'Dedicated Archive Room (Climate Controlled)',
                'secure_vault': 'Secure Storage Vault (Room 09)',
                'controlled_substance_storage': 'Secure Storage Vault (Room 09)'
            }
        }

    def _parse_operations_sheet(self, df: pd.DataFrame) -> Dict:
        """Parse Operations sheet."""
        values = df.set_index('Field')['Value'].to_dict()

        return {
            'scope': {
                'cultivation': values.get('Cultivation') == 'Yes',
                'flowering': values.get('Flowering') == 'Yes',
                'harvesting': values.get('Harvesting') == 'Yes',
                'drying_curing': values.get('Drying & Curing') == 'Yes',
                'processing': values.get('Processing') == 'Yes',
                'packaging': values.get('Packaging') == 'Yes',
                'quality_control': values.get('Quality Control') == 'Yes',
                'distribution': values.get('Distribution') == 'Yes'
            },
            'additional_operations': values.get('Additional Operations', ''),
            'excluded_operations': values.get('Excluded Operations', ''),
            'quality_objectives': {
                'objective_1': values.get('Objective 1', ''),
                'objective_2': values.get('Objective 2', ''),
                'objective_3': values.get('Objective 3', ''),
                'custom_objectives': []
            },
            'metrics': {
                'target_metric_1': values.get('Metric 1', ''),
                'target_metric_2': values.get('Metric 2', ''),
                'target_metric_3': values.get('Metric 3', '')
            }
        }

    def _parse_qc_sheet(self, df: pd.DataFrame) -> Dict:
        """Parse Quality Control sheet."""
        values = df.set_index('Field')['Value'].to_dict()

        return {
            'in_house_lab': values.get('In-House Lab') == 'Yes',
            'contract_lab': values.get('Contract Lab') == 'Yes',
            'contract_lab_details': {
                'name': values.get('Lab Name', ''),
                'address': values.get('Address', ''),
                'contact_name': values.get('Contact Name', ''),
                'phone': values.get('Phone', ''),
                'email': values.get('Email', ''),
                'accreditation': values.get('Accreditation', '')
            },
            'testing_capabilities': {
                'potency_testing': values.get('Potency Testing', ''),
                'microbial_testing': values.get('Microbial Testing', ''),
                'pesticide_testing': values.get('Pesticide Testing', ''),
                'heavy_metals_testing': values.get('Heavy Metals Testing', ''),
                'moisture_content': values.get('Moisture Content', ''),
                'visual_inspection': values.get('Visual Inspection', '')
            }
        }

    def _parse_doc_control_sheet(self, df: pd.DataFrame) -> Dict:
        """Parse Document Control sheet."""
        values = df.set_index('Field')['Value'].to_dict()

        locations = []
        for i in range(1, 4):
            loc = values.get(f'Location {i}', '')
            if loc:
                locations.append(loc)

        return {
            'retention_periods': {
                'batch_records': values.get('Batch Records', ''),
                'training_records': values.get('Training Records', ''),
                'equipment_qualification': values.get('Equipment Qualification', ''),
                'deviation_reports': values.get('Deviation Reports', ''),
                'audit_reports': values.get('Audit Reports', ''),
                'qms_documents': values.get('QMS Documents', '')
            },
            'controlled_locations': locations,
            'access_control': {
                'procedure': values.get('Procedure', ''),
                'backup_frequency': values.get('Backup Frequency', '')
            }
        }

    def _format_excel_workbook(self):
        """Apply formatting to Excel workbook."""
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = load_workbook(self.excel_path)

        # Format each sheet
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # Header row formatting
            header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
            header_font = Font(bold=True, color='FFFFFF', size=12)

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')

            # Section headers (cells starting with ===)
            section_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            section_font = Font(bold=True, size=11)

            for row in ws.iter_rows(min_row=2):
                cell = row[0]
                if cell.value and str(cell.value).startswith('==='):
                    for c in row:
                        c.fill = section_fill
                        c.font = section_font

            # Adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(self.excel_path)


def main():
    """Main function."""
    import argparse

    parser = argparse.ArgumentParser(description='Excel Integration Tool')
    parser.add_argument('--export', action='store_true', help='Export YAML to Excel')
    parser.add_argument('--import', dest='import_excel', action='store_true', help='Import Excel to YAML')
    args = parser.parse_args()

    base_dir = Path(__file__).parent.parent
    config_path = base_dir / 'config' / 'facility_data.yaml'
    excel_path = base_dir / 'config' / 'facility_data.xlsx'

    integrator = ExcelIntegration(config_path, excel_path)

    if args.export:
        success = integrator.yaml_to_excel()
        sys.exit(0 if success else 1)
    elif args.import_excel:
        success = integrator.excel_to_yaml()
        sys.exit(0 if success else 1)
    else:
        console.print("[yellow]Please specify --export or --import[/yellow]")
        parser.print_help()
        sys.exit(1)


if __name__ == "__main__":
    main()
