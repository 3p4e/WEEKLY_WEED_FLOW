#!/usr/bin/env python3
"""
QMS Document Fixer
Automatically fixes identified gaps in QMS documents based on the standard checklist.

Author: QMS Development Team
Date: 2026-01-16
"""

import json
import logging
import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("ERROR: python-docx not available. Install with: pip install python-docx")
    sys.exit(1)

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Standard header template
HEADER_TEMPLATE = """
PURELY PLANT DOOEL - QUALITY MANAGEMENT SYSTEM
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
Document No.: {doc_number}
Version: {version}
Effective Date: {effective_date}
Department: {department}
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
"""

# Standard section templates
SECTION_TEMPLATES = {
    "purpose": """
1. PURPOSE
This document establishes the standard operating procedure for {process_name}. 
The objective is to ensure consistent, compliant, and reproducible operations 
in accordance with EU GMP requirements.
""",
    
    "scope": """
2. SCOPE
This procedure applies to:
• All personnel involved in {process_name}
• All areas where this activity is performed
• All related equipment and materials

Exclusions: None unless otherwise specified.
""",
    
    "responsibilities": """
3. RESPONSIBILITIES

| Role | Responsibility |
|------|----------------|
| Production Manager | Overall implementation and compliance |
| QA Manager | Review, approval, and compliance monitoring |
| Operators | Execution according to this procedure |
| QC Personnel | Testing and verification as required |
""",
    
    "definitions": """
4. DEFINITIONS
• SOP: Standard Operating Procedure
• QA: Quality Assurance
• QC: Quality Control
• GMP: Good Manufacturing Practice
• EU GMP: European Union Good Manufacturing Practice (EudraLex Volume 4)
• CAPA: Corrective and Preventive Action
""",
    
    "references": """
5. REFERENCES
• EudraLex Volume 4 - EU Guidelines for Good Manufacturing Practice
• EudraLex Annex 1 - Manufacture of Sterile Medicinal Products
• EudraLex Annex 15 - Qualification and Validation
• ICH Q7 - Good Manufacturing Practice Guide for Active Pharmaceutical Ingredients
• Company Quality Manual (QAS-01-001)
""",
    
    "deviation": """
6. DEVIATION HANDLING
Any deviation from this procedure must be:
1. Immediately reported to the QA Manager
2. Documented using the Deviation Report Form (QAT-004)
3. Investigated within 24 hours
4. Corrective actions implemented as per CAPA procedure (QAS-02-001)

Reference: Master Deviation Handling SOP (QAS-02-001)
""",
    
    "revision_history": """
REVISION HISTORY
| Version | Date | Author | Description of Changes |
|---------|------|--------|------------------------|
| 1.0 | {date} | QA Team | Initial release |
""",
    
    "approval": """
DOCUMENT APPROVAL

Prepared by: _________________ Date: _________
             (Author Name/Title)

Reviewed by: _________________ Date: _________
             (Reviewer Name/Title)

Approved by: _________________ Date: _________
             (QA Manager)
"""
}


class QMSDocumentFixer:
    """Fixes identified gaps in QMS documents"""
    
    def __init__(self, gap_report_path: str):
        self.gap_report_path = Path(gap_report_path)
        self.gap_report = self._load_gap_report()
        self.fixed_count = 0
        self.error_count = 0
        self.backup_dir = Path("backups") / datetime.now().strftime("%Y%m%d_%H%M%S")
    
    def _load_gap_report(self) -> Dict[str, Any]:
        """Load the gap analysis report"""
        if not self.gap_report_path.exists():
            logger.error(f"Gap report not found: {self.gap_report_path}")
            return {"detailed_results": []}
        
        with open(self.gap_report_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        # Handle both list format (from sop_gap_report) and dict format (from qms_master)
        if isinstance(data, list):
            return {"detailed_results": data}
        return data
    
    def fix_all_documents(self, target_dir: str):
        """Fix all documents in the target directory based on gap report"""
        target_path = Path(target_dir)
        
        # Create backup directory
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        logger.info(f"Backups will be saved to: {self.backup_dir}")
        
        # Get all docx files
        docx_files = list(target_path.glob("**/*.docx"))
        logger.info(f"Found {len(docx_files)} .docx files to process")
        
        for docx_file in docx_files:
            self._fix_document(docx_file)
        
        logger.info(f"\n=== SUMMARY ===")
        logger.info(f"Fixed: {self.fixed_count} documents")
        logger.info(f"Errors: {self.error_count} documents")
    
    def _fix_document(self, doc_path: Path):
        """Fix a single document"""
        try:
            # Find gap info for this document
            gap_info = self._find_gap_info(doc_path.name)
            
            if not gap_info:
                logger.info(f"Skipping {doc_path.name} - no gap info found")
                return
            
            if gap_info.get("verdict") == "PASS":
                logger.info(f"Skipping {doc_path.name} - already passing")
                return
            
            gaps = gap_info.get("gaps", [])
            if not gaps:
                logger.info(f"Skipping {doc_path.name} - no gaps identified")
                return
            
            # Create backup
            backup_path = self.backup_dir / doc_path.name
            shutil.copy2(doc_path, backup_path)
            
            # Load document
            doc = Document(doc_path)
            
            # Apply fixes based on gaps
            changes_made = []
            
            for gap in gaps:
                if "Document Header" in gap:
                    self._add_header(doc, doc_path.name)
                    changes_made.append("Added header")
                
                elif "Purpose Section" in gap:
                    self._add_section(doc, "purpose", doc_path.stem)
                    changes_made.append("Added purpose")
                
                elif "Scope Section" in gap:
                    self._add_section(doc, "scope", doc_path.stem)
                    changes_made.append("Added scope")
                
                elif "Responsibilities Section" in gap:
                    self._add_section(doc, "responsibilities", doc_path.stem)
                    changes_made.append("Added responsibilities")
                
                elif "Definitions Section" in gap:
                    self._add_section(doc, "definitions", doc_path.stem)
                    changes_made.append("Added definitions")
                
                elif "GMP Compliance" in gap:
                    self._add_section(doc, "references", doc_path.stem)
                    changes_made.append("Added GMP references")
                
                elif "Deviation Handling" in gap:
                    self._add_section(doc, "deviation", doc_path.stem)
                    changes_made.append("Added deviation handling")
                
                elif "Revision History" in gap:
                    self._add_section(doc, "revision_history", doc_path.stem)
                    changes_made.append("Added revision history")
                
                elif "Approval Signatures" in gap:
                    self._add_section(doc, "approval", doc_path.stem)
                    changes_made.append("Added approval section")
            
            # Save document
            if changes_made:
                doc.save(doc_path)
                self.fixed_count += 1
                logger.info(f"✅ Fixed {doc_path.name}: {', '.join(changes_made)}")
            
        except Exception as e:
            self.error_count += 1
            logger.error(f"❌ Error fixing {doc_path.name}: {e}")
    
    def _find_gap_info(self, filename: str) -> Optional[Dict[str, Any]]:
        """Find gap info for a specific file"""
        for result in self.gap_report.get("detailed_results", []):
            if result.get("sop_file") == filename:
                return result
        return None
    
    def _add_header(self, doc: Document, filename: str):
        """Add standardized header to document"""
        # Extract document number from filename
        doc_number = filename.split("_")[0] if "_" in filename else "DOC-XXX"
        
        header_text = HEADER_TEMPLATE.format(
            doc_number=doc_number,
            version="1.0",
            effective_date=datetime.now().strftime("%Y-%m-%d"),
            department="Quality Assurance"
        )
        
        # Insert at beginning
        first_para = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
        
        # Create new paragraph at start
        new_para = doc.paragraphs[0].insert_paragraph_before(header_text)
        new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    def _add_section(self, doc: Document, section_type: str, process_name: str):
        """Add a section to the document"""
        template = SECTION_TEMPLATES.get(section_type, "")
        if not template:
            return
        
        # Format template
        content = template.format(
            process_name=process_name.replace("_", " "),
            date=datetime.now().strftime("%Y-%m-%d")
        )
        
        # Add to document
        doc.add_paragraph(content)


def main():
    # Paths
    project_root = "/home/azzu/PROJ/Cannabis EU GMP QMS Creator"
    gap_report_path = os.path.join(project_root, "data", "qms_master_analysis_report.json")
    target_dir = os.path.join(project_root, "00_MASTER_DOCUMENTS")
    
    print("\n" + "="*60)
    print("QMS DOCUMENT FIXER")
    print("="*60 + "\n")
    
    fixer = QMSDocumentFixer(gap_report_path)
    fixer.fix_all_documents(target_dir)
    
    print(f"\nBackups saved to: {fixer.backup_dir}")


if __name__ == "__main__":
    main()
