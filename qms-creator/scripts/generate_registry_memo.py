import yaml
from docx import Document
from pathlib import Path
import subprocess

def load_registry():
    with open('config/document_registry.yaml', 'r') as f:
        return yaml.safe_load(f)

def format_registry_content(registry):
    lines = []
    for fam_id, fam_data in registry['families'].items():
        fam_name_en = fam_data.get('name_en', 'Unknown')
        fam_name_mk = fam_data.get('name_mk', 'Непознато')
        lines.append(f"\n### {fam_name_mk} | {fam_name_en} ({fam_id})")
        
        sops = fam_data.get('sops', {})
        for sop_id, sop_data in sops.items():
            title_en = sop_data.get('title_en', 'Untitled')
            title_mk = sop_data.get('title_mk', 'Без наслов')
            lines.append(f"- **{sop_id}**: {title_mk} | {title_en}")
            
            annexes = sop_data.get('annexes', {})
            for ann_id, ann_data in annexes.items():
                a_title_en = ann_data.get('title_en', 'Untitled')
                a_title_mk = ann_data.get('title_mk', 'Без наслов')
                lines.append(f"  - *{ann_id}*: {a_title_mk} | {a_title_en}")
    return "\n".join(lines)

def replace_text_in_paragraph(paragraph, old_text, new_text):
    if old_text in paragraph.text:
        # Simple replacement
        paragraph.text = paragraph.text.replace(old_text, new_text)

def replace_text_in_document(doc, replacements):
    for paragraph in doc.paragraphs:
        for old_text, new_text in replacements.items():
            replace_text_in_paragraph(paragraph, old_text, new_text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for old_text, new_text in replacements.items():
                        replace_text_in_paragraph(paragraph, old_text, new_text)

def main():
    registry = load_registry()
    registry_text = format_registry_content(registry)
    
    src_path = Path("01_QUALITY_ASSURANCE/QA_00.04_A01_General_Memorandum_v1.0_EN.docx")
    dest_docx = Path("01_QUALITY_ASSURANCE/QA_00.04_MEM_Document_Registry_v1.0_EN.docx")
    
    if not src_path.exists():
        print(f"Error: Template not found at {src_path}")
        return

    doc = Document(src_path)
    
    replacements = {
        "MEMYY_DD_nnn": "MEM25_QA_REG",
        "___ / ___ / 20__": "15 / 01 / 2025",
        "[Замени со Име, Организација, Оддел | Replace with Name / Organization / Department: ]": "All Departments | Сите оддели",
        "[Замени со Име, Организација, Оддел | Replace with Name / Organization / Department:]": "Quality Assurance | Осигурување на квалитет",
        "[ SUBJECT:  Replace with Document Type by format and intended use, ex. REQUEST / COMPLAINT / NOTATION / URGENT INSTRUCTION etc.]": "Purely Plant Document Registry | Регистар на документи на Purely Plant",
        "[CONTENT: Replace with intended content in bilingual formatting “MKD text | Eng Text”]": registry_text
    }
    
    replace_text_in_document(doc, replacements)
    doc.save(dest_docx)
    print(f"Saved {dest_docx}")
    
    # Convert to PDF
    try:
        subprocess.run([
            "soffice", "--headless", "--convert-to", "pdf", 
            "--outdir", "01_QUALITY_ASSURANCE", str(dest_docx)
        ], check=True)
        print(f"Generated PDF for {dest_docx}")
    except Exception as e:
        print(f"Error generating PDF: {e}")

if __name__ == "__main__":
    main()
