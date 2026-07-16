import yaml
from docx import Document
from pathlib import Path
import subprocess

def load_registry():
    with open('config/document_registry.yaml', 'r') as f:
        return yaml.safe_load(f)

def format_registry_content(registry):
    lines = []
    for fam_id, fam_data in registry['families'].items():
        fam_name_en = fam_data.get('name_en', 'Unknown')
        fam_name_mk = fam_data.get('name_mk', 'Непознато')
        lines.append(f"\n### {fam_name_mk} | {fam_name_en} ({fam_id})")
        sops = fam_data.get('sops', {})
        for sop_id, sop_data in sops.items():
            title_en = sop_data.get('title_en', 'Untitled')
            title_mk = sop_data.get('title_mk', 'Без наслов')
            lines.append(f"- **{sop_id}**: {title_mk} | {title_en}")
            annexes = sop_data.get('annexes', {})
            for ann_id, ann_data in annexes.items():
                a_title_en = ann_data.get('title_en', 'Untitled')
                a_title_mk = ann_data.get('title_mk', 'Без наслов')
                lines.append(f"  - *{ann_id}*: {a_title_mk} | {a_title_en}")
    return "\n".join(lines)

def update_header(doc, doc_id, title):
    for section in doc.sections:
        header = section.header
        if header:
            for table in header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if '[Document Type and Description]' in cell.text:
                            cell.text = cell.text.replace('[Document Type and Description]', title)
                        if 'eff. dd.mm.yy' in cell.text:
                            cell.text = cell.text.replace('eff. dd.mm.yy', '15.01.25')
                        if 'QMS-DC-SOP-XXX' in cell.text:
                            cell.text = cell.text.replace('QMS-DC-SOP-XXX', doc_id)

def replace_text_in_document(doc, replacements):
    for paragraph in doc.paragraphs:
        for old_text, new_text in replacements.items():
            if old_text in paragraph.text:
                paragraph.text = paragraph.text.replace(old_text, new_text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for old_text, new_text in replacements.items():
                        if old_text in paragraph.text:
                            paragraph.text = paragraph.text.replace(old_text, new_text)

def main():
    registry = load_registry()
    registry_text = format_registry_content(registry)
    
    memo_id = "MEM25_QA_001"
    memo_title = "Purely Plant QMS Document Registry Release"
    
    # --- 1. GENERATE MEMO (A01) ---
    doc = Document("01_QUALITY_ASSURANCE/QA_00.04_A01_General_Memorandum_v1.0_EN.docx")
    update_header(doc, "QA_00.04_A01_v1", memo_title)
    
    replacements = {
        "MEMYY_DD_nnn": memo_id,
        "___ / ___ / 20__": "15 / 01 / 2025",
        "[Замени со Име, Организација, Оддел | Replace with Name / Organization / Department: ]": "All Department Heads | Сите раководители на оддели",
        "[Замени со Име, Организација, Оддел | Replace with Name / Organization / Department:]": "Quality Assurance | Осигурување на квалитет",
        "[ SUBJECT:  Replace with Document Type by format and intended use, ex. REQUEST / COMPLAINT / NOTATION / URGENT INSTRUCTION etc.]": "OFFICIAL RELEASE: Purely Plant Document Registry",
        "[CONTENT: Replace with intended content in bilingual formatting “MKD text | Eng Text”]": "The following registry defines the official document hierarchy for Purely Plant GmbH. All departments must align their local documentation to these codes.\n" + registry_text
    }
    replace_text_in_document(doc, replacements)
    doc.save("01_QUALITY_ASSURANCE/QA_00.04_MEM_Document_Registry_v1.0_EN.docx")

    # --- 2. GENERATE DISTRIBUTION RECORD (A02) ---
    doc_dist = Document("01_QUALITY_ASSURANCE/QA_00.04_A02_Distribution_Record_v1.0_EN.docx")
    update_header(doc_dist, "QA_00.04_A02_v1", "Distribution Record: " + memo_id)
    
    dist_replacements = {
        "MEM____/____/____": memo_id,
        "Subject of Memorandum": memo_title,
        "Issuing Department": "Quality Assurance",
        "Date of Distribution": "15.01.2025"
    }
    replace_text_in_document(doc_dist, dist_replacements)
    
    depts = ["Cultivation", "Production", "QC", "QA", "Maintenance", "Logistics", "HR", "Security"]
    for table in doc_dist.tables:
        if len(table.rows) > 0 and len(table.rows[0].cells) >= 5:
            # Check text in the first row cells manually
            header_text = " ".join([c.text for c in table.rows[0].cells])
            if "NAME" in header_text or "#" in header_text:
                for i in range(1, 26):
                    if i < len(table.rows):
                        row = table.rows[i]
                        row.cells[1].text = f"Dept Head / Representative {i}"
                        row.cells[2].text = depts[i % len(depts)]
    
    doc_dist.save("01_QUALITY_ASSURANCE/QA_00.04_A02_REG_Distribution_Record.docx")

    # --- 3. GENERATE ISSUANCE LIST (A03) ---
    doc_iss = Document("01_QUALITY_ASSURANCE/QA_00.04_A03_Memorandum_Issuance_List_v1.0_EN.docx")
    update_header(doc_iss, "QA_00.04_A03_v1", "Memorandum Issuance Log")
    
    iss_replacements = {
        "MEM_         _      /": "MEM_QA_01/25",
        "Issuing Department": "Quality Assurance",
        "Date Issued": "15.01.2025",
        "MEMYY_DD_001": "MEM25_QA_001"
    }
    replace_text_in_document(doc_iss, iss_replacements)
    
    for table in doc_iss.tables:
        if len(table.rows) > 0 and len(table.rows[0].cells) >= 9:
            header_text = " ".join([c.text for c in table.rows[0].cells])
            if "REF. NO." in header_text:
                row = table.rows[1]
                row.cells[1].text = memo_id
                row.cells[2].text = "15.01.25"
                row.cells[3].text = "All Depts"
                row.cells[4].text = "QMS Registry Release"
                row.cells[5].text = "AD"
                row.cells[6].text = "Y"
                row.cells[7].text = "Open"
                row.cells[8].text = "C"
                break
    
    doc_iss.save("01_QUALITY_ASSURANCE/QA_00.04_A03_REG_Issuance_Log.docx")

    # Convert all to PDF
    for f in ["QA_00.04_MEM_Document_Registry_v1.0_EN.docx", "QA_00.04_A02_REG_Distribution_Record.docx", "QA_00.04_A03_REG_Issuance_Log.docx"]:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", "01_QUALITY_ASSURANCE", "01_QUALITY_ASSURANCE/" + f])

if __name__ == "__main__":
    main()
