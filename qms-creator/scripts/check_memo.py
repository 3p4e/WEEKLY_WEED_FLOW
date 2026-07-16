#!/usr/bin/env python3
"""
Script to extract the document registry table from the generated memo document.
This is needed because the table generation logic in generate_registry_package.py
might be inserting the table into a paragraph instead of replacing the content properly.
"""

import sys
from pathlib import Path

from docx import Document


def extract_table_content(docx_path):
    """Extract and print table content from a DOCX file."""
    try:
        doc = Document(docx_path)
        print(f"File: {docx_path}")
        print(f"Total tables: {len(doc.tables)}")

        for i, table in enumerate(doc.tables):
            print(f"\nTable {i}: {len(table.rows)} rows x {len(table.columns)} columns")

            # Print first few rows
            for j, row in enumerate(table.rows[:5]):
                cells = [cell.text.strip() for cell in row.cells]
                print(f"  Row {j}: {cells}")

    except Exception as e:
        print(f"Error reading {docx_path}: {e}")


def main():
    memo_path = Path("01_QUALITY_ASSURANCE/QA_00.04_MEM_Registry_Issuance_v1.0_EN.docx")
    if memo_path.exists():
        extract_table_content(memo_path)
    else:
        print(f"File not found: {memo_path}")


if __name__ == "__main__":
    main()
