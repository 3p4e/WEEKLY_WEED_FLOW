#!/usr/bin/env python3
"""
SOP Gap Analyzer
Compares existing SOPs against target standard checklist.

Author: QMS Development Team
Date: 2025-01-22
"""

import json
import logging
import os
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

import yaml

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class SOPGapAnalyzer:
    """Analyzes SOPs against a standard checklist"""
    
    def __init__(self, checklist_path: str):
        self.checklist_path = Path(checklist_path)
        self.checklist = self._load_checklist()
    
    def _load_checklist(self) -> Dict[str, Any]:
        """Load the standard checklist from YAML"""
        if not self.checklist_path.exists():
            logger.error(f"Checklist not found: {self.checklist_path}")
            return {}
        
        with open(self.checklist_path, 'r', encoding='utf-8') as f:
            return yaml.safe_load(f)
    
    def analyze_sop(self, sop_path: str) -> Dict[str, Any]:
        """Analyze a single SOP against the checklist"""
        sop_path = Path(sop_path)
        
        if not sop_path.exists():
            return {"error": f"SOP file not found: {sop_path}"}
        
        # Extract content
        content = self._extract_content(sop_path)
        if not content:
            return {"error": "Could not extract content from SOP"}
        
        # Analyze sections
        section_results = self._check_sections(content)
        
        # Analyze quality checks
        quality_results = self._check_quality(content)
        
        # Calculate score
        total_score, max_score = self._calculate_score(section_results, quality_results)
        percentage = round((total_score / max_score) * 100) if max_score > 0 else 0
        
        # Determine verdict
        threshold_pass = self.checklist.get("scoring", {}).get("threshold_pass", 70)
        threshold_conditional = self.checklist.get("scoring", {}).get("threshold_conditional", 50)
        
        if percentage >= threshold_pass:
            verdict = "PASS"
        elif percentage >= threshold_conditional:
            verdict = "CONDITIONAL PASS"
        else:
            verdict = "NEEDS IMPROVEMENT"
        
        return {
            "sop_file": sop_path.name,
            "score": percentage,
            "verdict": verdict,
            "sections": section_results,
            "quality_checks": quality_results,
            "gaps": self._extract_gaps(section_results, quality_results),
            "strengths": self._extract_strengths(section_results, quality_results)
        }
    
    def _extract_content(self, sop_path: Path) -> str:
        """Extract text content from SOP file"""
        suffix = sop_path.suffix.lower()
        
        if suffix == '.docx':
            if not DOCX_AVAILABLE:
                logger.warning("python-docx not available")
                return ""
            try:
                doc = Document(sop_path)
                full_text = []
                
                # Add headers
                for section in doc.sections:
                    for p in section.header.paragraphs:
                        full_text.append(p.text)
                
                # Add body
                for p in doc.paragraphs:
                    full_text.append(p.text)
                
                # Add footers
                for section in doc.sections:
                    for p in section.footer.paragraphs:
                        full_text.append(p.text)
                        
                return "\n".join(full_text)
            except Exception as e:
                logger.error(f"Error reading docx: {e}")
                return ""
        
        elif suffix in ['.txt', '.md']:
            try:
                with open(sop_path, 'r', encoding='utf-8') as f:
                    return f.read()
            except Exception as e:
                logger.error(f"Error reading text file: {e}")
                return ""
        
        else:
            logger.warning(f"Unsupported file type: {suffix}")
            return ""
    
    def _check_sections(self, content: str) -> List[Dict[str, Any]]:
        """Check for required sections"""
        results = []
        content_lower = content.lower()
        
        for section in self.checklist.get("required_sections", []):
            section_id = section.get("id", "unknown")
            section_name = section.get("name", section_id)
            keywords = section.get("keywords", [])
            weight = section.get("weight", 5)
            
            # Check if section exists
            found = False
            matched_keywords = []
            
            for keyword in keywords:
                if keyword.lower() in content_lower:
                    found = True
                    matched_keywords.append(keyword)
            
            # Additional heuristics
            if section_id == "procedure":
                # Check for numbered steps
                if re.search(r'\d+\.\s+\w', content):
                    found = True
                    matched_keywords.append("numbered steps")
            
            results.append({
                "section_id": section_id,
                "section_name": section_name,
                "found": found,
                "matched_keywords": matched_keywords,
                "weight": weight,
                "score": weight if found else 0
            })
        
        return results
    
    def _check_quality(self, content: str) -> List[Dict[str, Any]]:
        """Check quality criteria"""
        results = []
        content_lower = content.lower()
        
        for check in self.checklist.get("quality_checks", []):
            check_id = check.get("id", "unknown")
            check_name = check.get("name", check_id)
            patterns = check.get("patterns", [])
            weight = check.get("weight", 5)
            
            found = False
            matched_patterns = []
            
            for pattern in patterns:
                if pattern.lower() in content_lower:
                    found = True
                    matched_patterns.append(pattern)
            
            results.append({
                "check_id": check_id,
                "check_name": check_name,
                "found": found,
                "matched_patterns": matched_patterns,
                "weight": weight,
                "score": weight if found else 0
            })
        
        return results
    
    def _calculate_score(self, sections: List[Dict], quality: List[Dict]) -> tuple:
        """Calculate total and max scores"""
        total = sum(s["score"] for s in sections) + sum(q["score"] for q in quality)
        max_score = sum(s["weight"] for s in sections) + sum(q["weight"] for q in quality)
        return total, max_score
    
    def _extract_gaps(self, sections: List[Dict], quality: List[Dict]) -> List[str]:
        """Extract list of gaps"""
        gaps = []
        
        for s in sections:
            if not s["found"]:
                gaps.append(f"Missing section: {s['section_name']}")
        
        for q in quality:
            if not q["found"]:
                gaps.append(f"Quality check failed: {q['check_name']}")
        
        return gaps
    
    def _extract_strengths(self, sections: List[Dict], quality: List[Dict]) -> List[str]:
        """Extract list of strengths"""
        strengths = []
        
        for s in sections:
            if s["found"]:
                strengths.append(f"Has {s['section_name']}")
        
        for q in quality:
            if q["found"]:
                strengths.append(f"Includes {q['check_name']}")
        
        return strengths[:5]  # Limit to top 5


def analyze_folder(folder_path: str, checklist_path: str) -> List[Dict[str, Any]]:
    """Analyze all SOPs in a folder"""
    analyzer = SOPGapAnalyzer(checklist_path)
    folder = Path(folder_path)
    
    results = []
    for sop_file in folder.glob("**/*.docx"):
        logger.info(f"Analyzing: {sop_file.name}")
        result = analyzer.analyze_sop(str(sop_file))
        results.append(result)
    
    return results


if __name__ == "__main__":
    # Default paths
    checklist_path = "/home/azzu/PROJ/Cannabis EU GMP QMS Creator/CONTENT_CREATOR_FRAMEWORK/sop_standard_checklist.yaml"
    sop_folder = "/home/azzu/PROJ/Cannabis EU GMP QMS Creator/REFERENCE_MATERIALS/PP/Purely Plant InUse SOPs"
    
    # Run analysis
    results = analyze_folder(sop_folder, checklist_path)
    
    # Print summary
    print("\n=== SOP Gap Analysis Summary ===")
    for r in results:
        if "error" in r:
            print(f"❌ {r.get('sop_file', 'unknown')}: {r['error']}")
        else:
            icon = "✅" if r["verdict"] == "PASS" else "⚠️" if r["verdict"] == "CONDITIONAL PASS" else "❌"
            print(f"{icon} {r['sop_file']}: {r['score']}% - {r['verdict']}")
            if r["gaps"]:
                for gap in r["gaps"][:3]:
                    print(f"   → {gap}")
    
    # Save full report
    output_path = "/home/azzu/PROJ/Cannabis EU GMP QMS Creator/data/sop_gap_report.json"
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(results, f, indent=2, ensure_ascii=False)
    print(f"\nFull report saved to: {output_path}")
