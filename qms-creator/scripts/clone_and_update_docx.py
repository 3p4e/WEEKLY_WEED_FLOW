import os
import sys
from pathlib import Path
from docx import Document

def replace_text_in_paragraph(paragraph, old_text, new_text):
    if old_text in paragraph.text:
        inline = paragraph.runs
        # Simple replacement for now. 
        # Note: This might break formatting if the text is split across runs.
        # A more robust way is to replace the text in the full paragraph.text and reset runs, 
        # but that loses formatting.
        # For this task, we'll try a simple replace on the full text if possible, 
        # or just iterate runs if the text is contained within a single run.
        
        # Strategy: Reconstruct text, do replace, then put back? No, formatting loss.
        # Strategy: Check each run.
        for run in inline:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text)
        
        # Fallback: if text is split across runs, this won't work.
        # But for IDs like "QMS-DC-SOP-XXX", they are likely in one run.

def replace_text_in_document(doc, replacements):
    # Body paragraphs
    for paragraph in doc.paragraphs:
        for old_text, new_text in replacements.items():
            replace_text_in_paragraph(paragraph, old_text, new_text)

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for old_text, new_text in replacements.items():
                        replace_text_in_paragraph(paragraph, old_text, new_text)
                # Nested tables?
                for nested_table in cell.tables:
                    for nrow in nested_table.rows:
                        for ncell in nrow.cells:
                            for nparagraph in ncell.paragraphs:
                                for old_text, new_text in replacements.items():
                                    replace_text_in_paragraph(nparagraph, old_text, new_text)

    # Headers and Footers
    for section in doc.sections:
        # Header
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header:
                for paragraph in header.paragraphs:
                    for old_text, new_text in replacements.items():
                        replace_text_in_paragraph(paragraph, old_text, new_text)
                for table in header.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                for old_text, new_text in replacements.items():
                                    replace_text_in_paragraph(paragraph, old_text, new_text)
        # Footer
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer:
                for paragraph in footer.paragraphs:
                    for old_text, new_text in replacements.items():
                        replace_text_in_paragraph(paragraph, old_text, new_text)
                for table in footer.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                for old_text, new_text in replacements.items():
                                    replace_text_in_paragraph(paragraph, old_text, new_text)

def main():
    base_dir = Path("REFERENCE_MATERIALS/Memo SOP")
    target_dir = Path("01_QUALITY_ASSURANCE")
    target_dir.mkdir(exist_ok=True)

    mappings = [
        (
            "QMS-DC-SOP-XXX_Memorandum_Control_Management_V2.0.docx",
            "QA_00.04_Memorandum_Control_Management_v1.0_EN.docx"
        ),
        (
            "QMS-DC-SOP-XXX-ANX-01_General_Memorandum_v4_EDITED.docx",
            "QA_00.04_A01_General_Memorandum_v1.0_EN.docx"
        ),
        (
            "QMS-DC-SOP-XXX-ANX-02_Distribution_Record_v4.docx",
            "QA_00.04_A02_Distribution_Record_v1.0_EN.docx"
        ),
        (
            "QMS-DC-SOP-XXX-ANX-03_Memorandum_Issuance_List_v5.docx",
            "QA_00.04_A03_Memorandum_Issuance_List_v1.0_EN.docx"
        ),
        (
            "QMS-DC-SOP-XXX-ANX-04_Memorandum_Receipt_List_v2.docx",
            "QA_00.04_A04_Memorandum_Receipt_List_v1.0_EN.docx"
        )
    ]

    replacements = {
        "QMS-DC-SOP-XXX-ANX-01": "QA_00.04_A01_v1",
        "QMS-DC-SOP-XXX-ANX-02": "QA_00.04_A02_v1",
        "QMS-DC-SOP-XXX-ANX-03": "QA_00.04_A03_v1",
        "QMS-DC-SOP-XXX-ANX-04": "QA_00.04_A04_v1",
        "QMS-DC-SOP-XXX": "QA_00.04_v1",
        "Effective Date: [Date]": "Effective Date: 15.01.25",
        "Date: [Date]": "Date: 15.01.25",
        "v4": "v1.0",
        "V2.0": "v1.0",
        "v5": "v1.0",
        "v2": "v1.0"
    }

    for src_name, dest_name in mappings:
        src_path = base_dir / src_name
        dest_path = target_dir / dest_name
        
        if not src_path.exists():
            print(f"Warning: Source file not found: {src_path}")
            continue

        print(f"Processing {src_name} -> {dest_name}")
        doc = Document(src_path)
        replace_text_in_document(doc, replacements)
        doc.save(dest_path)
        print(f"Saved {dest_path}")

if __name__ == "__main__":
    main()
