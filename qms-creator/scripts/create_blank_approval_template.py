#!/usr/bin/env python3
"""
Create Blank Approval Page Template in DOCX Format
Generates a Word document with only the approval page structure
"""

from pathlib import Path
from typing import Optional
import sys
from datetime import datetime, timedelta

sys.path.insert(0, str(Path(__file__).parent))

from facility_config import load_facility_config

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("Error: python-docx not installed. Install with: pip install python-docx")
    sys.exit(1)


class BlankApprovalTemplateGenerator:
    """Generate blank approval page template in DOCX format."""

    def __init__(self, config=None):
        """Initialize template generator."""
        if config is None:
            self.config = load_facility_config()
        else:
            self.config = config

        self.branding = self.config.get('document_control.branding', {})
        self.approval_chain = self.config.get('document_control.approval_chain', {})
        self.storage = self.config.get('document_control.storage', {})
        self.revision_policy = self.config.get('document_control.revision_policy', {})

    def add_table_borders(self, table, color='70ad47', width='18'):
        """Add borders to table with specified color."""
        tbl = table._element
        tblPr = tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement('w:tblPr')
            tbl.insert(0, tblPr)

        tblBorders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), width)
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), color)
            tblBorders.append(border)

        tblPr.append(tblBorders)

    def shade_cell(self, cell, fill_color='e2efd9'):
        """Add background color to cell."""
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), fill_color)
        cell._element.get_or_add_tcPr().append(shading_elm)

    def generate_template(self, output_path: Path):
        """Generate blank approval page template in DOCX format."""
        doc = Document()

        # Set document margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # ==================== HEADER TABLE ====================
        header_table = doc.add_table(rows=1, cols=2)
        header_table.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.add_table_borders(header_table)

        # Header cells
        left_cell = header_table.rows[0].cells[0]
        right_cell = header_table.rows[0].cells[1]

        # Shade header
        self.shade_cell(left_cell, 'e2efd9')
        self.shade_cell(right_cell, 'e2efd9')

        # Left cell - Document ID
        left_para = left_cell.paragraphs[0]
        left_para.text = 'N/A'
        left_para.runs[0].bold = True
        left_para.runs[0].font.size = Pt(10)

        # Right cell - Page number
        right_para = right_cell.paragraphs[0]
        right_para.text = 'Page N/A of N/A'
        right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right_para.runs[0].font.size = Pt(10)

        doc.add_paragraph()

        # ==================== TITLE SECTION ====================
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run('Purely Plant GmbH\n')
        title_run.bold = True
        title_run.font.size = Pt(14)

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle.add_run('[DOCUMENT TITLE - MACEDONIAN] | [DOCUMENT TITLE - ENGLISH]')
        subtitle_run.font.size = Pt(12)
        subtitle_run.italic = True

        doc.add_paragraph()

        # ==================== DOCUMENT INFO ====================
        info_para = doc.add_paragraph()
        info_para.add_run('Document ID: ').bold = True
        info_para.add_run('N/A  |  ')
        info_para.add_run('Version: ').bold = True
        info_para.add_run('N/A  |  ')
        info_para.add_run('Effective Date: ').bold = True
        info_para.add_run('N/A')
        info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # ==================== APPROVAL TABLE ====================
        approval_heading = doc.add_paragraph('ОДОБРУВАЊЕ НА ДОКУМЕНТ | DOCUMENT APPROVAL')
        approval_heading.runs[0].bold = True
        approval_heading.runs[0].font.size = Pt(11)

        approval_table = doc.add_table(rows=4, cols=5)
        approval_table.style = 'Light Grid Accent 1'
        self.add_table_borders(approval_table)

        # Header row
        header_cells = approval_table.rows[0].cells
        headers = ['Дејство / Action', 'Позиција / Position', 'Име и Презиме / Name', 'Датум / Date', 'Потпис / Signature']

        for idx, header_text in enumerate(headers):
            cell = header_cells[idx]
            self.shade_cell(cell, 'e2efd9')
            para = cell.paragraphs[0]
            para.text = header_text
            para.runs[0].bold = True
            para.runs[0].font.size = Pt(9)

        # Data rows
        actions = ['Подготвено од: / Prepared by:', 'Проверено од: / Checked by:', 'Одобрено од: / Approved by:']

        for row_idx, action in enumerate(actions):
            cells = approval_table.rows[row_idx + 1].cells
            cells[0].text = action
            cells[1].text = 'N/A'
            cells[2].text = 'N/A'
            cells[3].text = '__________'
            cells[4].text = '_________________'

            for cell in cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)

        doc.add_paragraph()

        # ==================== DOCUMENT CONTROL TABLE ====================
        control_heading = doc.add_paragraph('КОНТРОЛА НА ДОКУМЕНТ | DOCUMENT CONTROL')
        control_heading.runs[0].bold = True
        control_heading.runs[0].font.size = Pt(11)

        control_table = doc.add_table(rows=5, cols=2)
        self.add_table_borders(control_table)

        control_items = [
            ('Оригиналот се чува во: / The original is kept & stored in:', 'N/A'),
            ('Тип на документ: / Type of document:', 'N/A'),
            ('Овој документ се чува во: / This document is kept & stored in:', 'N/A'),
            ('Тип на копија: / Copy type:', 'N/A'),
            ('Корисник / User:', '☐ Внатрешен / Internal  ☐ Надворешен / External'),
        ]

        for row_idx, (label, value) in enumerate(control_items):
            row = control_table.rows[row_idx]
            row.cells[0].text = label
            row.cells[1].text = value

            # Shade label cell
            self.shade_cell(row.cells[0], 'e2efd9')

            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)

        # Add revision info
        doc.add_paragraph()
        rev_para = doc.add_paragraph()
        rev_para.add_run('Датум на последна ревизија: / Date of last revision: ').bold = True
        rev_para.add_run('__________  |  ')
        rev_para.add_run('Ефективен датум: / Effective Date: ').bold = True
        rev_para.add_run('__________  |  ')
        rev_para.add_run('Датум на валидност: / Validity Date: ').bold = True
        rev_para.add_run('__________')

        for run in rev_para.runs:
            run.font.size = Pt(9)

        doc.add_paragraph()

        # ==================== REVISION HISTORY ====================
        history_heading = doc.add_paragraph('ИСТОРИЈА НА РЕВИЗИИ | REVISION HISTORY')
        history_heading.runs[0].bold = True
        history_heading.runs[0].font.size = Pt(11)

        history_table = doc.add_table(rows=4, cols=3)
        self.add_table_borders(history_table)

        history_headers = ['Верзија бр. / Version No.', 'Датум / Date of change', 'Опис / Description of change']

        # Header row
        for idx, header_text in enumerate(history_table.rows[0].cells):
            self.shade_cell(header_text, 'e2efd9')
            para = header_text.paragraphs[0]
            para.text = history_headers[idx]
            para.runs[0].bold = True
            para.runs[0].font.size = Pt(9)

        # Data rows
        for row_idx in range(1, 4):
            row = history_table.rows[row_idx]
            row.cells[0].text = 'N/A'
            row.cells[1].text = 'N/A'
            row.cells[2].text = 'N/A'

            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)

        doc.add_paragraph()
        doc.add_paragraph('_' * 80)
        doc.add_paragraph()

        # Save document
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path


def main():
    """Main function to generate blank template."""
    import argparse

    parser = argparse.ArgumentParser(description='Generate blank approval page template')
    parser.add_argument(
        '--output',
        type=str,
        default='BLANK_APPROVAL_PAGE_TEMPLATE_v1.0.docx',
        help='Output file path'
    )

    args = parser.parse_args()

    base_dir = Path(__file__).parent.parent
    output_path = base_dir / args.output

    print("\n" + "=" * 80)
    print("GENERATING BLANK APPROVAL PAGE TEMPLATE")
    print("=" * 80 + "\n")

    config = load_facility_config()
    generator = BlankApprovalTemplateGenerator(config)

    try:
        result = generator.generate_template(output_path)
        print(f"✅ Template created successfully!")
        print(f"📄 File: {result}")
        print(f"📊 Size: {result.stat().st_size / 1024:.1f} KB\n")
        print("How to use:")
        print("  1. Open the document in Microsoft Word or compatible software")
        print("  2. Replace 'N/A' values with document-specific information")
        print("  3. Complete signature fields with initials and dates")
        print("  4. Save as new document with your naming convention")
        print(f"\n✨ Template ready to use!\n")

    except Exception as e:
        print(f"❌ Error generating template: {str(e)}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
